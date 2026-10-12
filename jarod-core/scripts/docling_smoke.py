"""Test de fumée B20 : profil docling-slim SANS PyTorch, hors ligne, sur des documents FABRIQUÉS ici (aucun document personnel)."""
import os, sys, socket, tempfile, traceback, json

os.environ["HF_HUB_OFFLINE"] = "1"
os.environ["TRANSFORMERS_OFFLINE"] = "1"
attempts = []
_orig = socket.socket.connect
def _blocked(self, addr):
    attempts.append(str(addr))
    raise OSError("réseau interdit (test de fumée)")
socket.socket.connect = _blocked
socket.socket.connect_ex = lambda self, addr: (attempts.append(str(addr)), 1)[1]

report = {}
try:
    import torch  # noqa
    report["import_torch"] = "RÉUSSI (inattendu)"
except ImportError as e:
    report["import_torch"] = "échoue comme prévu : " + type(e).__name__

d = tempfile.mkdtemp(prefix="jarod-smoke-")

# --- documents fabriqués ---
import docx, pptx, openpyxl
from docx.shared import Pt
doc = docx.Document()
doc.add_heading("Rapport de test", 1)
doc.add_paragraph("Premier paragraphe en français avec des accents : éàçù.")
doc.add_heading("Section deux", 2)
t = doc.add_table(rows=2, cols=2)
t.cell(0, 0).text = "Nom"; t.cell(0, 1).text = "Valeur"; t.cell(1, 0).text = "alpha"; t.cell(1, 1).text = "42"
doc.save(os.path.join(d, "t.docx"))

prs = pptx.Presentation()
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = "Titre de diapositive"
s.placeholders[1].text = "Un point important\nUn second point"
prs.save(os.path.join(d, "t.pptx"))

wb = openpyxl.Workbook()
ws = wb.active; ws.title = "Ventes"
ws.append(["Produit", "Quantité"]); ws.append(["pommes", 12]); ws.append(["poires", 7])
wb.create_sheet("Notes").append(["remarque", "ok"])
wb.save(os.path.join(d, "t.xlsx"))

open(os.path.join(d, "t.html"), "w", encoding="utf8").write(
    "<html><body><h1>Page HTML</h1><p>Un paragraphe.</p><h2>Liste</h2><ul><li>un</li><li>deux</li></ul>"
    "<table><tr><th>a</th><th>b</th></tr><tr><td>1</td><td>2</td></tr></table></body></html>")

def make_pdf(path, pages):
    objs = []
    n = 3 + 2 * len(pages)  # 1 catalogue, 2 pages, 3 police, puis (page, contenu)*
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(len(pages)))
    objs.append("<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>")
    objs.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, lines in enumerate(pages):
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>")
        stream = "BT /F1 12 Tf 72 720 Td 16 TL " + " T* ".join(f"({l}) Tj" for l in lines) + " ET"
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    out = b"%PDF-1.4\n"
    offs = []
    for i, o in enumerate(objs):
        offs.append(len(out))
        out += f"{i + 1} 0 obj\n{o}\nendobj\n".encode("latin-1")
    x = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for o in offs:
        out += f"{o:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{x}\n%%EOF\n".encode()
    open(path, "wb").write(out)
make_pdf(os.path.join(d, "t.pdf"), [["Titre du PDF", "Premiere page : du texte simple."], ["Deuxieme page : autre texte."]])

# --- conversion ---
def summarize(res):
    doc = res.document
    md = doc.export_to_markdown()
    return dict(status=str(res.status), pages=len(doc.pages), md_chars=len(md), headings=sum(1 for i, _ in doc.iterate_items() if type(i).__name__ == "SectionHeaderItem"),
                tables=len(doc.tables), md_head=md[:160].replace("\n", " | "))

try:
    from docling.document_converter import DocumentConverter, PdfFormatOption
    from docling.datamodel.base_models import InputFormat
    report["import_docling"] = "ok"
except Exception as e:
    report["import_docling"] = "ÉCHEC : " + repr(e)
    print(json.dumps(report, ensure_ascii=False, indent=1)); sys.exit(2)

conv = {}
try:
    c = DocumentConverter()
    for ext in ("docx", "pptx", "xlsx", "html"):
        try:
            conv[ext] = summarize(c.convert(os.path.join(d, "t." + ext)))
        except Exception as e:
            conv[ext] = "ÉCHEC : " + repr(e)[:300]
except Exception as e:
    conv["converter"] = "ÉCHEC : " + repr(e)[:300]

try:
    from docling.pipeline.native_pdf_pipeline import NativePdfPipeline
    from docling.datamodel.pipeline_options import NativePdfPipelineOptions
    c2 = DocumentConverter(format_options={InputFormat.PDF: PdfFormatOption(pipeline_cls=NativePdfPipeline, pipeline_options=NativePdfPipelineOptions())})
    conv["pdf"] = summarize(c2.convert(os.path.join(d, "t.pdf")))
except Exception as e:
    conv["pdf"] = "ÉCHEC : " + repr(e)[:400] + " | " + traceback.format_exc()[-400:]

report["conversions"] = conv
report["torch_charge"] = "torch" in sys.modules
report["tentatives_reseau"] = attempts
print(json.dumps(report, ensure_ascii=False, indent=1))
