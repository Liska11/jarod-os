import docx, pptx, openpyxl
d = docx.Document(); d.add_heading("Titre Un", 1); d.add_paragraph("Un paragraphe de test."); t = d.add_table(rows=2, cols=2); t.cell(0,0).text="a"; t.cell(0,1).text="b"; t.cell(1,0).text="1"; t.cell(1,1).text="2"; d.save("t.docx")
p = pptx.Presentation(); s = p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text="Diapo"; s.placeholders[1].text="puce un"; p.save("t.pptx")
w = openpyxl.Workbook(); ws = w.active; ws.title="Feuille"; ws.append(["x","y"]); ws.append([1,2]); w.save("t.xlsx")
open("t.html","w").write("<h1>Page</h1><p>Texte.</p>")
