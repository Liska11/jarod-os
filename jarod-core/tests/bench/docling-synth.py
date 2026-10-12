#!/usr/bin/env python3
"""
Corpus SYNTHÉTIQUE de la gate B (jeu de validation, DECISIONS 2026-09-30). Déterministe : mêmes entrées => mêmes contenus.
Source : les documents du dépôt (plan, gap-reel, PHASE-3). Sortie : <out>/ (ignoré par git, reproductible).
Il contient EXPRÈS ce que le jeu de mise au point a montré difficile : titres, tableaux, cellules fusionnées, liens
hypertexte, dates, nombres, puces, PDF à couche texte. Généré avec python-docx / python-pptx / openpyxl / fpdf,
lu ensuite par Docling ; la référence, elle, relit le XML brut (docling-reference.py) — elle ne passe pas par ces bibliothèques.
Usage : python docling-synth.py <racine du dépôt> <dossier de sortie>
"""
import random
import re
import sys
from datetime import datetime, timedelta
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(sys.argv[1])
OUT = Path(sys.argv[2])
OUT.mkdir(parents=True, exist_ok=True)
URL = "https://exemple.invalid/jarod/"


def clean(s: str) -> str:
    s = re.sub(r"\*\*|__|`", "", s)
    return re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", s).strip()


def blocks(path: Path):
    """Markdown => [("h", niveau, texte) | ("t", [[cellules]]) | ("b", texte) | ("p", texte)]."""
    out, table = [], []
    for raw in path.read_text(encoding="utf8").splitlines():
        line = raw.rstrip()
        if line.startswith("|"):
            cells = [clean(c) for c in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c.replace(" ", "")) for c in cells if c):
                continue
            table.append(cells)
            continue
        if table:
            out.append(("t", table))
            table = []
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            out.append(("h", len(m.group(1)), clean(m.group(2))))
        elif re.match(r"^\s*[-*]\s+", line):
            out.append(("b", clean(re.sub(r"^\s*[-*]\s+", "", line))))
        elif line.strip():
            out.append(("p", clean(line)))
    if table:
        out.append(("t", table))
    return out


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    rid = part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


def make_docx(src: Path, dest: Path, limit: int):
    doc = Document()
    n = 0
    for b in blocks(src):
        if n >= limit:
            break
        n += 1
        if b[0] == "h":
            doc.add_heading(b[2], level=min(b[1], 4))
            if b[1] == 2:
                p = doc.add_paragraph("Référence : ")
                add_hyperlink(p, f"fiche {n}", f"{URL}fiche-{n}?a=%20b")
        elif b[0] == "b":
            doc.add_paragraph(b[1], style="List Bullet")
        elif b[0] == "p":
            doc.add_paragraph(b[1])
        else:
            rows = b[1]
            cols = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=cols)
            t.style = "Table Grid"
            for i, r in enumerate(rows):
                for j in range(cols):
                    t.cell(i, j).text = r[j] if j < len(r) else ""
            if len(rows) >= 3 and cols >= 3:
                t.cell(len(rows) - 1, 0).merge(t.cell(len(rows) - 1, 1))  # cellule fusionnée
    doc.save(dest)


def make_pptx(src: Path, dest: Path, slides: int):
    prs = Presentation()
    count = 0
    cur_title, cur = None, []

    def flush():
        nonlocal count
        if cur_title is None or count >= slides:
            return
        count += 1
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = cur_title
        body = s.placeholders[1].text_frame
        body.clear()
        items = [x for x in cur if x[0] in ("b", "p")][:6]
        for k, x in enumerate(items):
            (body.paragraphs[0] if k == 0 else body.add_paragraph()).text = x[1][:160]
        tabs = [x for x in cur if x[0] == "t"]
        if tabs:
            rows = tabs[0][1][:5]
            cols = min(max(len(r) for r in rows), 4)
            shape = s.shapes.add_table(len(rows), cols, Inches(0.5), Inches(4.3), Inches(9), Inches(1.6))
            for i, r in enumerate(rows):
                for j in range(cols):
                    shape.table.cell(i, j).text = (r[j] if j < len(r) else "")[:40]

    for b in blocks(src):
        if b[0] == "h" and b[1] <= 2:
            flush()
            cur_title, cur = b[2][:80], []
        else:
            cur.append(b)
    flush()
    prs.save(dest)


def make_xlsx(src: Path, dest: Path, sheets: int, seed: int):
    rng = random.Random(seed)
    wb = Workbook()
    wb.remove(wb.active)
    k = 0
    for b in blocks(src):
        if b[0] != "t" or k >= sheets:
            continue
        k += 1
        ws = wb.create_sheet(f"Feuille{k}")
        rows = b[1]
        cols = max(len(r) for r in rows)
        ws["A1"] = f"Tableau {k} — {rows[0][0][:30]}"
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(cols, 2))  # titre fusionné
        for i, r in enumerate(rows, start=3):
            for j in range(cols):
                ws.cell(i, j + 1, r[j] if j < len(r) else "")
        for c in ws[3]:
            c.font = Font(bold=True)
        # Colonnes chiffrées et datées, tirées d'un générateur à graine fixe.
        base = datetime(2026, 3, 1, 8, 0, 0)
        ws.cell(3, cols + 1, "Quantité")
        ws.cell(3, cols + 2, "Échéance")
        for i in range(4, 3 + len(rows)):
            ws.cell(i, cols + 1, rng.randint(1, 500))
            d = ws.cell(i, cols + 2, base + timedelta(days=rng.randint(0, 200), hours=rng.randint(0, 9)))
            d.number_format = "yyyy-mm-dd hh:mm"
    wb.save(dest)


LATIN = {"—": "-", "–": "-", "’": "'", "‘": "'", "“": '"', "”": '"', "…": "...", "•": "-", "→": "->", "≥": ">=", "≤": "<=", "×": "x", "✅": "OK", "⚠": "!", " ": " "}


def latin1(s: str) -> str:
    for a, b in LATIN.items():
        s = s.replace(a, b)
    return s.encode("latin-1", "replace").decode("latin-1")


def make_pdf(src: Path, dest: Path, limit: int):
    pdf = FPDF()
    pdf.set_auto_page_break(True, 15)
    pdf.add_page()
    n = 0
    for b in blocks(src):
        if n >= limit:
            break
        n += 1
        if b[0] == "h":
            pdf.set_font("Helvetica", "B", 14 - min(b[1], 4))
            pdf.multi_cell(0, 7, latin1(b[2]), new_x="LMARGIN", new_y="NEXT")
        elif b[0] in ("p", "b"):
            pdf.set_font("Helvetica", "", 10)
            pdf.multi_cell(0, 5, latin1(("- " if b[0] == "b" else "") + b[1]), new_x="LMARGIN", new_y="NEXT")
        else:
            pdf.set_font("Courier", "", 8)
            for r in b[1]:
                pdf.multi_cell(0, 4, latin1(" | ".join(r))[:120], new_x="LMARGIN", new_y="NEXT")
        pdf.ln(1)
    pdf.output(str(dest))


plan, gap, ph3 = ROOT / "plan.md", ROOT / "gap-reel.md", ROOT / "PHASE-3.md"
make_docx(plan, OUT / "synth-01-plan.docx", 140)
make_docx(gap, OUT / "synth-02-gap.docx", 140)
make_docx(ph3, OUT / "synth-03-phase3.docx", 160)
make_pptx(plan, OUT / "synth-04-plan.pptx", 14)
make_pptx(gap, OUT / "synth-05-gap.pptx", 14)
make_xlsx(ph3, OUT / "synth-06-phase3.xlsx", 8, seed=1)
make_xlsx(plan, OUT / "synth-07-plan.xlsx", 8, seed=2)
make_pdf(plan, OUT / "synth-08-plan.pdf", 260)
make_pdf(gap, OUT / "synth-09-gap.pdf", 260)
print("synthétique prêt :", len(list(OUT.iterdir())), "fichiers")
